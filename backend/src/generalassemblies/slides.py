"""Lecture et remplissage du PowerPoint d'une assemblee generale.

L'appli ne decide rien du contenu : l'IA (voir ai.py) dit quel texte mettre
dans quelle zone de quelle diapo, et ce module le place dans le PPT modele
sans toucher a la mise en page (meme police, meme taille, meme niveau de
puce que le texte d'origine). "À compléter" est surligne en jaune, comme
dans les PPT du tresorier.

Vignettes des diapos : PPT converti en PDF par LibreOffice (soffice), puis
chaque page en image PNG (pymupdf). Sans LibreOffice, l'ecran affiche
seulement les textes des diapos.
"""

from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
import tempfile
import threading
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

TO_COMPLETE = "À compléter"
HIGHLIGHT = "FFFF00"
# Elements de a:rPr qui doivent venir apres a:highlight (schema OOXML).
_AFTER_HIGHLIGHT = [
    qn(f"a:{tag}")
    for tag in ("uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst")
]
_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
THUMBNAIL_WIDTH = 960
_thumbnail_lock = threading.Lock()


class SlidesError(ValueError):
    """PPT illisible ou modification impossible (message affichable)."""


def open_presentation(path: str | os.PathLike):
    try:
        return Presentation(path if hasattr(path, "read") else str(path))
    except Exception as exc:  # noqa: BLE001 - tout fichier illisible
        raise SlidesError("Ce fichier n'est pas un PowerPoint (.pptx) lisible") from exc


def _text_zones(shapes):
    """Zones de texte modifiables d'une diapo (slide.shapes) : formes qui
    contiennent deja du texte, y compris dans les groupes (les zones vides
    servent souvent de fond a des images ou des icones : on n'y ecrit pas).
    Les SmartArt ne sont pas modifiables."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _text_zones(shape.shapes)
        elif shape.has_text_frame and shape.text_frame.text.strip():
            yield shape


def _is_title(shape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in _TITLE_TYPES


def _level(paragraph_xml) -> int:
    ppr = paragraph_xml.find(qn("a:pPr"))
    return int(ppr.get("lvl", "0")) if ppr is not None else 0


def _hidden(slide) -> bool:
    """Diapo masquee dans PowerPoint (pas projetee, pas dans le PDF)."""
    return slide._element.get("show") == "0"


def outline(path: str | os.PathLike) -> list[dict]:
    """Plan du PPT, envoye a l'IA : pour chaque diapo, ses zones de texte
    (identifiant, titre ou texte, paragraphes avec leur niveau de puce).
    Noms de champs en francais : ils sont lus par l'IA."""
    presentation = open_presentation(path)
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        zones = []
        for shape in _text_zones(slide.shapes):
            zones.append(
                {
                    "zone": shape.shape_id,
                    "role": "titre" if _is_title(shape) else "texte",
                    "paragraphes": [
                        {"niveau": p.level, "texte": p.text.replace("\v", "\n")}
                        for p in shape.text_frame.paragraphs
                        if p.text.strip()
                    ],
                }
            )
        slides.append({"diapo": number, "masquee": _hidden(slide), "zones": zones})
    return slides


def texts(path: str | os.PathLike) -> list[dict]:
    """Textes de chaque diapo, pour l'apercu a l'ecran : {"number", "hidden",
    "title", "lines": [{"level", "text"}]}."""
    presentation = open_presentation(path)
    result = []
    for number, slide in enumerate(presentation.slides, start=1):
        title, lines = "", []
        for shape in _text_zones(slide.shapes):
            if _is_title(shape) and not title:
                # Symboles de police (Wingdings...) : fleche lisible a l'ecran.
                title = re.sub("[\uf000-\uf0ff]", "→", shape.text_frame.text.replace("\v", " ")).strip()
                continue
            lines += [
                {"level": p.level, "text": p.text.replace("\v", " ").strip()}
                for p in shape.text_frame.paragraphs
                if p.text.strip()
            ]
        result.append({"number": number, "hidden": _hidden(slide), "title": title, "lines": lines})
    return result


def _set_highlight(run_xml) -> None:
    rpr = run_xml.find(qn("a:rPr"))
    if rpr is None:
        rpr = run_xml.makeelement(qn("a:rPr"), {})
        run_xml.insert(0, rpr)
    for old in rpr.findall(qn("a:highlight")):
        rpr.remove(old)
    highlight = rpr.makeelement(qn("a:highlight"), {})
    color = highlight.makeelement(qn("a:srgbClr"), {"val": HIGHLIGHT})
    highlight.append(color)
    following = next((child for child in rpr if child.tag in _AFTER_HIGHLIGHT), None)
    if following is None:
        rpr.append(highlight)
    else:
        following.addprevious(highlight)


def _new_paragraph(model, level: int, text: str):
    """Paragraphe au format de model (copie de son XML), avec le texte text.
    "À compléter" est mis dans un morceau a part, surligne en jaune."""
    paragraph = copy.deepcopy(model)
    runs = paragraph.findall(qn("a:r"))
    run_model = copy.deepcopy(runs[0]) if runs else None
    for child in list(paragraph):
        if child.tag in (qn("a:r"), qn("a:br"), qn("a:fld")):
            paragraph.remove(child)
    if run_model is not None and run_model.find(qn("a:rPr")) is not None:
        for old in run_model.find(qn("a:rPr")).findall(qn("a:highlight")):
            old.getparent().remove(old)

    ppr = paragraph.find(qn("a:pPr"))
    if level or ppr is not None:
        if ppr is None:
            ppr = paragraph.makeelement(qn("a:pPr"), {})
            paragraph.insert(0, ppr)
        if level:
            ppr.set("lvl", str(level))
        elif "lvl" in ppr.attrib:
            del ppr.attrib["lvl"]

    end = paragraph.find(qn("a:endParaRPr"))

    def add(element) -> None:
        if end is None:
            paragraph.append(element)
        else:
            end.addprevious(element)

    # Retour a la ligne dans le paragraphe ("\n") : saut de ligne a:br.
    for line_index, line in enumerate(text.split("\n")):
        if line_index:
            br = paragraph.makeelement(qn("a:br"), {})
            if run_model is not None and run_model.find(qn("a:rPr")) is not None:
                br.append(copy.deepcopy(run_model.find(qn("a:rPr"))))
            add(br)
        for index, piece in enumerate(re.split(f"({re.escape(TO_COMPLETE)})", line)):
            if not piece:
                continue
            if run_model is not None:
                run = copy.deepcopy(run_model)
                run.find(qn("a:t")).text = piece
            else:
                run = paragraph.makeelement(qn("a:r"), {})
                t = run.makeelement(qn("a:t"), {})
                t.text = piece
                run.append(t)
            if index % 2 == 1:
                _set_highlight(run)
            add(run)
    return paragraph


def _replace_text(shape, paragraphs: list[dict]) -> None:
    body = shape.text_frame._txBody
    originals = body.findall(qn("a:p"))
    with_text = [p for p in originals if "".join(t.text or "" for t in p.iter(qn("a:t"))).strip()] or originals
    news = []
    for item in paragraphs:
        level = max(0, min(int(item.get("niveau", 0)), 4))
        # Format d'un paragraphe d'origine du meme niveau, sinon le plus proche.
        model = min(with_text, key=lambda p: (abs(_level(p) - level), with_text.index(p)))
        news.append(_new_paragraph(model, level, str(item.get("texte", "")).strip()))
    for paragraph in originals:
        body.remove(paragraph)
    for paragraph in news or [_new_paragraph(with_text[0], 0, "")]:
        body.append(paragraph)


def fill(template: str | os.PathLike, output: str | os.PathLike, zones: list[dict]) -> list[int]:
    """Ecrit dans output le PPT template ou chaque zone listee recoit ses
    nouveaux paragraphes : [{"diapo", "zone", "paragraphes": [{"niveau",
    "texte"}]}]. Zones inconnues ignorees. Retourne les diapos modifiees."""
    presentation = open_presentation(template)
    slides = list(presentation.slides)
    changed = set()
    for item in zones:
        number = item.get("diapo")
        if not isinstance(number, int) or not 1 <= number <= len(slides):
            continue
        shape = next((s for s in _text_zones(slides[number - 1].shapes) if s.shape_id == item.get("zone")), None)
        if shape is None or not item.get("paragraphes"):
            continue
        _replace_text(shape, item["paragraphes"])
        changed.add(number)
    presentation.save(str(output))
    return sorted(changed)


# ----- garde-fou sur les montants -----

_AMOUNT = re.compile(r"(?<![\d,.])(\d{1,3}(?:[   .]\d{3})+|\d+)(?:,(\d{1,2}))?\s*(?:€|euros?\b)")


def amounts_in(text: str) -> list[float]:
    """Montants en euros d'un texte : "15 718 €", "4 794,36 €", "+1507€"."""
    values = []
    for whole, cents in _AMOUNT.findall(text):
        value = float(re.sub(r"[   .]", "", whole))
        if cents:
            value += int(cents.ljust(2, "0")) / 100
        values.append(value)
    return values


def check_amounts(zones: list[dict], allowed: set[float]) -> list[str]:
    """Montants ecrits par l'IA qui ne sont pas dans les donnees fournies
    (a l'euro pres) : une alerte par diapo et par montant."""
    known = {abs(round(a)) for a in allowed if a is not None}
    warnings = []
    for item in zones:
        for paragraph in item.get("paragraphes", []):
            for value in amounts_in(str(paragraph.get("texte", ""))):
                if not any(abs(round(value) - k) <= 1 for k in known):
                    amount = f"{round(value):,}".replace(",", " ")
                    message = f"Diapo {item.get('diapo')} : {amount} € ne vient pas du bilan, à vérifier."
                    if message not in warnings:
                        warnings.append(message)
    return warnings


# ----- vignettes -----


def soffice() -> str | None:
    return os.environ.get("LIBREOFFICE_PATH") or shutil.which("soffice") or shutil.which("libreoffice")


def thumbnails(pptx: str | os.PathLike, folder: str | os.PathLike) -> list[Path] | None:
    """Images PNG des diapos affichees (folder/slide-1.png... : les diapos
    masquees ne sont pas dans le PDF de LibreOffice), creees une seule fois.
    None si LibreOffice n'est pas installe ou si la conversion echoue."""
    folder = Path(folder)
    done = folder / "done"
    with _thumbnail_lock:
        if not done.exists():
            program = soffice()
            if program is None:
                return None
            import pymupdf

            with tempfile.TemporaryDirectory() as tmp:
                try:
                    subprocess.run(
                        [
                            program,
                            f"-env:UserInstallation=file://{tmp}/profile",
                            "--headless",
                            "--convert-to",
                            "pdf",
                            "--outdir",
                            tmp,
                            str(pptx),
                        ],
                        check=True,
                        capture_output=True,
                        timeout=180,
                    )
                    pdf = Path(tmp) / (Path(pptx).stem + ".pdf")
                    folder.mkdir(parents=True, exist_ok=True)
                    with pymupdf.open(pdf) as document:
                        for number, page in enumerate(document, start=1):
                            zoom = THUMBNAIL_WIDTH / page.rect.width
                            page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(str(folder / f"slide-{number}.png"))
                except (subprocess.SubprocessError, OSError, RuntimeError, ValueError):
                    return None
            done.write_text("ok")
    return sorted(folder.glob("slide-*.png"), key=lambda p: int(p.stem.split("-")[1]))
