"""Telechargement d'un bilan financier calcule : Excel (.xlsx, openpyxl) ou
PDF (pymupdf), avec le meme contenu que l'ecran : verification des soldes,
recettes et depenses par categorie comparees a la saison precedente,
analyse, compte d'exploitation par mois et, dans l'Excel, toutes les
operations de la saison."""

from __future__ import annotations

import html
import io

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONTHS_FR = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]


def _month(month: str) -> str:
    """"2025-07" -> "juil.-25" (comme le modele du tresorier)."""
    return f"{MONTHS_FR[int(month[5:7]) - 1]}-{month[2:4]}"


def _fr(iso: str) -> str:
    year, month, day = iso.split("-")
    return f"{day}/{month}/{year}"


def _eur(value: float | None) -> str:
    if value is None:
        return "—"
    text = f"{abs(value):,.2f}".replace(",", " ").replace(".", ",")
    return ("−" if value < 0 else "") + text + " €"


def _pct(pct: int | None) -> str:
    return "" if pct is None else f"{pct:+d} %"


def comparison(result: dict, side: str) -> list[dict]:
    """Lignes recettes (side "income") ou depenses ("expense") : categorie,
    montant de la saison, de la saison precedente et ecart en %."""
    previous = (result.get("previous") or {}).get("net", {})
    sign = 1 if side == "income" else -1
    rows = []
    for row in result[side]:
        before = previous.get(row["category"])
        before = None if before is None else sign * before
        pct = None if not before else round((row["amount"] - before) / abs(before) * 100)
        rows.append({"category": row["category"], "amount": row["amount"], "previous": before, "pct": pct})
    return rows


def _title(report: dict) -> str:
    return f"Bilan financier {report['result']['seasonName']} : {report['name']}"


def _verification_text(result: dict) -> str:
    v = result["verification"]
    if result["opening"] is None or result["closing"] is None:
        return "Vérification impossible : soldes inconnus (relevés manquants)."
    text = (
        f"Solde au {_fr(result['start'])} {_eur(result['opening'])} + recettes {_eur(result['totalIncome'])} "
        f"− dépenses {_eur(result['totalExpense'])} = {_eur(v['expected'])} ; solde réel au "
        f"{_fr(result['cutoff'])} : {_eur(v['actual'])}."
    )
    return text + (" Vérifié, écart 0,00 €." if v["ok"] else f" Écart : {_eur(v['gap'])}.")


def to_xlsx(report: dict) -> bytes:
    result = report["result"]
    previous_name = (result.get("previous") or {}).get("season")
    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E9E3F5")
    money = '#,##0.00 "€";-#,##0.00 "€";""'
    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Bilan"
    sheet.append([_title(report)])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Du {_fr(result['start'])} au {_fr(result['cutoff'])}" + ("" if result["cutoff"] == result["end"] else " (saison en cours)")])
    sheet.append([f"Calculé par {result['modelLabel']} le {_fr(result['generatedAt'][:10])}"])
    sheet.append([])
    sheet.append([_verification_text(result)])
    sheet.append([])
    for side, label in (("income", "Recettes"), ("expense", "Dépenses")):
        sheet.append([label, result["seasonName"], previous_name or "", "Écart"])
        for cell in sheet[sheet.max_row]:
            cell.font, cell.fill = bold, head
        for row in comparison(result, side):
            sheet.append([row["category"], row["amount"], row["previous"], None if row["pct"] is None else row["pct"] / 100])
            sheet.cell(sheet.max_row, 2).number_format = money
            sheet.cell(sheet.max_row, 3).number_format = money
            sheet.cell(sheet.max_row, 4).number_format = "+0%;-0%;0%"
        total = result["totalIncome"] if side == "income" else result["totalExpense"]
        sheet.append([f"Total {label.lower()}", total])
        sheet.cell(sheet.max_row, 1).font = bold
        sheet.cell(sheet.max_row, 2).font = bold
        sheet.cell(sheet.max_row, 2).number_format = money
        sheet.append([])
    sheet.append(["Résultat", result["result"]])
    sheet.cell(sheet.max_row, 1).font = bold
    sheet.cell(sheet.max_row, 2).number_format = money
    sheet.append([])
    sheet.append(["Analyse"])
    sheet.cell(sheet.max_row, 1).font = bold
    for line in result["analysis"]:
        sheet.append([line])
    sheet.column_dimensions["A"].width = 34
    for col in "BCD":
        sheet.column_dimensions[col].width = 16

    months = workbook.create_sheet("Par mois")
    categories = result["categories"]
    months.append([f"COMPTE D'EXPLOITATION {result['seasonName']} ÉTABLI SUIVANT LES RELEVÉS DE BANQUE"])
    months["A1"].font = bold
    months.append(["Mois", "Solde au 1er", *categories, "Solde fin de mois"])
    for cell in months[2]:
        cell.font, cell.fill = bold, head
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    for m in result["months"]:
        months.append([_month(m["month"]), m["opening"], *[m["amounts"].get(c) for c in categories], m["closing"]])
    totals = [sum(m["amounts"].get(c) or 0 for m in result["months"]) for c in categories]
    months.append(["Vérif", result["opening"], *totals, result["verification"]["expected"]])
    for cell in months[months.max_row]:
        cell.font = bold
    for row in months.iter_rows(min_row=3, min_col=2):
        for cell in row:
            cell.number_format = money
    months.column_dimensions["A"].width = 10
    for i in range(2, len(categories) + 4):
        months.column_dimensions[get_column_letter(i)].width = 13

    operations = workbook.create_sheet("Opérations")
    operations.append(["Date", "Libellé", "Détails", "Montant", "Catégorie", "Classée par l'IA"])
    for cell in operations[1]:
        cell.font, cell.fill = bold, head
    for op in result["operations"]:
        operations.append(
            [_fr(op["date"]), op["label"], op["details"], op["amount"], op["category"], "oui" if op["category"] != op["bankCategory"] else ""]
        )
        operations.cell(operations.max_row, 4).number_format = money
    for col, width in zip("ABCDEF", (11, 40, 40, 13, 26, 14)):
        operations.column_dimensions[col].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


_CSS = """
* { font-family: sans-serif; }
body { font-size: 9pt; color: #222; }
h1 { font-size: 15pt; margin: 0 0 2pt; }
h2 { font-size: 11pt; margin: 10pt 0 4pt; }
p { margin: 2pt 0; }
.muted { color: #666; }
.check { padding: 5pt; border: 1px solid #9ab; background: #eef4fa; }
.alert { padding: 5pt; border: 1px solid #d96; background: #fdf0e6; }
table { border-collapse: collapse; width: 100%; }
th { background: #ece6f6; text-align: right; padding: 2pt 3pt; font-weight: bold; }
td { text-align: right; padding: 2pt 3pt; border-bottom: 1px solid #ddd; }
th:first-child, td:first-child { text-align: left; }
tr.total td { font-weight: bold; border-top: 1px solid #555; }
.small td, .small th { font-size: 6.5pt; padding: 1pt 2pt; }
"""


def _html(report: dict) -> str:
    result = report["result"]
    e = html.escape
    previous_name = (result.get("previous") or {}).get("season")
    period = f"Du {_fr(result['start'])} au {_fr(result['cutoff'])}" + ("" if result["cutoff"] == result["end"] else " (saison en cours)")
    parts = [
        f"<h1>{e(_title(report))}</h1>",
        f"<p class='muted'>{period}. Calculé par {e(result['modelLabel'])} le {_fr(result['generatedAt'][:10])}.</p>",
        f"<p class='{'check' if result['verification']['ok'] else 'alert'}'>{e(_verification_text(result))}</p>",
    ]
    for side, label in (("income", "Recettes"), ("expense", "Dépenses")):
        total = result["totalIncome"] if side == "income" else result["totalExpense"]
        rows = "".join(
            f"<tr><td>{e(r['category'])}</td><td>{_eur(r['amount'])}</td>"
            + (f"<td>{_eur(r['previous'])}</td><td>{_pct(r['pct'])}</td>" if previous_name else "")
            + "</tr>"
            for r in comparison(result, side)
        )
        header = f"<th>{label}</th><th>{e(result['seasonName'])}</th>" + (f"<th>{e(previous_name)}</th><th>Écart</th>" if previous_name else "")
        parts.append(
            f"<h2>{label}</h2><table><tr>{header}</tr>{rows}"
            f"<tr class='total'><td>Total</td><td>{_eur(total)}</td>{'<td></td><td></td>' if previous_name else ''}</tr></table>"
        )
    parts.append(f"<h2>Résultat : {_eur(result['result'])}</h2>")
    parts.append("<h2>Analyse</h2><ol>" + "".join(f"<li>{e(line)}</li>" for line in result["analysis"]) + "</ol>")
    categories = result["categories"]
    head = "".join(f"<th>{e(c)}</th>" for c in ["Mois", "Solde au 1er", *categories, "Solde fin"])
    body = "".join(
        f"<tr><td>{_month(m['month'])}</td><td>{_eur(m['opening'])}</td>"
        + "".join(f"<td>{_eur(m['amounts'][c]) if c in m['amounts'] else ''}</td>" for c in categories)
        + f"<td>{_eur(m['closing'])}</td></tr>"
        for m in result["months"]
    )
    totals = "".join(f"<td>{_eur(sum(m['amounts'].get(c) or 0 for m in result['months']))}</td>" for c in categories)
    parts.append(
        f"<h2>Compte d'exploitation par mois</h2><table class='small'><tr>{head}</tr>{body}"
        f"<tr class='total'><td>Vérif</td><td>{_eur(result['opening'])}</td>{totals}<td>{_eur(result['verification']['expected'])}</td></tr></table>"
    )
    return "".join(parts)


def to_pdf(report: dict) -> bytes:
    buffer = io.BytesIO()
    page = pymupdf.paper_rect("a4-l")
    where = page + (36, 36, -36, -36)
    story = pymupdf.Story(html=_html(report), user_css=_CSS)
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    return buffer.getvalue()
